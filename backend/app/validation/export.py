import io
import json
from typing import List, Dict, Any
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.validation.models import ValidationSummary, QuestionValidation


def export_validation_report_rows(questions: List[Any]) -> List[Dict[str, Any]]:
    """Transforms a list of questions with attached validation results into structured report rows."""
    rows = []
    for q in questions:
        q_num = getattr(q, "question_number", 1)
        page_num = getattr(q, "page_number", 1)
        val = getattr(q, "validation", None)

        if isinstance(val, dict):
            status = val.get("status", "PERFECT")
            score = val.get("score", 100.0)
            errors = val.get("errors", [])
        elif isinstance(val, QuestionValidation):
            status = val.status
            score = val.score
            errors = [e.model_dump() if hasattr(e, "model_dump") else e for e in val.errors]
        else:
            status = "PERFECT"
            score = 100.0
            errors = []

        err_codes = set()
        err_msgs = []
        for e in errors:
            c = e.get("code", "") if isinstance(e, dict) else getattr(e, "code", "")
            m = e.get("message", "") if isinstance(e, dict) else getattr(e, "message", "")
            if c:
                err_codes.add(c)
            if m:
                err_msgs.append(f"[{c}] {m}")

        # Missing & empty option details
        missing_opts = []
        empty_opts = []
        for e in errors:
            c = e.get("code", "") if isinstance(e, dict) else getattr(e, "code", "")
            f = e.get("field", "") if isinstance(e, dict) else getattr(e, "field", "")
            if c == "MISSING_OPTION":
                missing_opts.append(f.replace("options.", ""))
            elif c == "EMPTY_OPTION":
                empty_opts.append(f.replace("options.", ""))

        latex_err_types = [c for c in err_codes if "LATEX" in c]
        image_crop_err_types = [c for c in err_codes if "CROP" in c]

        row = {
            "Question Number": q_num,
            "Page": page_num,
            "Status": status,
            "Quality Score": score,
            "Question Error": "YES" if any(c.startswith("QUESTION_") for c in err_codes) else "NO",
            "MCQ Error": "YES" if any(c in ("MISSING_OPTION", "EMPTY_OPTION", "DUPLICATE_OPTION", "INVALID_ANSWER", "OPTION_ORDER_ERROR", "OPTION_MIXING_ERROR") for c in err_codes) else "NO",
            "Missing Options": ", ".join(missing_opts) if missing_opts else "None",
            "Empty Options": ", ".join(empty_opts) if empty_opts else "None",
            "Duplicate Options": "YES" if "DUPLICATE_OPTION" in err_codes else "NO",
            "Invalid Answer": "YES" if "INVALID_ANSWER" in err_codes else "NO",
            "LaTeX Error": "YES" if any("LATEX" in c for c in err_codes) else "NO",
            "LaTeX Error Type": ", ".join(latex_err_types) if latex_err_types else "None",
            "Image Expected": "YES" if bool(getattr(q, "has_question_image", False) or getattr(q, "image_url", None) or "IMAGE_MISSING" in err_codes) else "NO",
            "Image Missing": "YES" if "IMAGE_MISSING" in err_codes else "NO",
            "Image Extraction Error": "YES" if "IMAGE_EXTRACTION_ERROR" in err_codes else "NO",
            "Image Crop Error": ", ".join(image_crop_err_types) if image_crop_err_types else "NO",
            "OCR Warning": "YES" if "OCR_SUSPICIOUS" in err_codes else "NO",
            "Duplicate Question": "YES" if ("DUPLICATE_QUESTION" in err_codes or "POSSIBLE_DUPLICATE_QUESTION" in err_codes) else "NO",
            "Error Details": "\n".join(err_msgs) if err_msgs else "No errors detected (Perfect extraction).",
        }
        rows.append(row)
    return rows


def export_validation_to_excel_bytes(questions: List[Any], summary: Optional[ValidationSummary] = None) -> bytes:
    """Generates a beautifully styled Microsoft Excel (.xlsx) validation report."""
    rows = export_validation_report_rows(questions)
    df = pd.DataFrame(rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Validation Report"

    # Header style
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # Status fills
    fill_perfect = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # Green
    fill_warning = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Yellow
    fill_fault = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")    # Red
    fill_critical = PatternFill(start_color="FCE7F3", end_color="FCE7F3", fill_type="solid") # Rose

    # Write Headers
    columns = list(df.columns)
    ws.append(columns)
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = border_thin

    # Write Data
    for r_idx, row_data in enumerate(rows, start=2):
        row_vals = [row_data[c] for c in columns]
        ws.append(row_vals)

        status_val = str(row_data.get("Status", "PERFECT")).upper()
        current_status_fill = None
        if status_val == "PERFECT":
            current_status_fill = fill_perfect
        elif status_val == "WARNING":
            current_status_fill = fill_warning
        elif status_val == "FAULT":
            current_status_fill = fill_fault
        elif status_val == "CRITICAL":
            current_status_fill = fill_critical

        for c_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = border_thin
            cell.alignment = align_left if columns[c_idx - 1] == "Error Details" else align_center
            if columns[c_idx - 1] == "Status" and current_status_fill:
                cell.fill = current_status_fill
                cell.font = Font(name="Calibri", size=10, bold=True)

    # Auto adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        if col_letter == get_column_letter(len(columns)):  # Error details column
            ws.column_dimensions[col_letter].width = 50
        else:
            ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 30))

    ws.row_dimensions[1].height = 28

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_validation_to_csv_bytes(questions: List[Any]) -> bytes:
    """Generates UTF-8 encoded CSV bytes with BOM for Excel compatibility."""
    rows = export_validation_report_rows(questions)
    df = pd.DataFrame(rows)
    csv_str = df.to_csv(index=False, encoding="utf-8-sig")
    return csv_str.encode("utf-8-sig")


def export_validation_to_json_bytes(questions: List[Any], summary: Optional[ValidationSummary] = None) -> bytes:
    """Generates structured JSON validation report bytes."""
    data = {
        "summary": summary.model_dump() if summary else {},
        "questions": export_validation_report_rows(questions),
    }
    return json.dumps(data, indent=2, ensure_ascii=False).encode("utf-8")
